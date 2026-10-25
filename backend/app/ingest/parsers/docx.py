"""DOCX text extraction: python-docx first, raw `word/document.xml` as fallback."""

from __future__ import annotations

import io
import re
import zipfile
import zlib

from app.ingest.parsers import Extracted, ExtractError
from app.ingest.text import normalize

_TAG = re.compile(r"<[^>]+>")
_PARA_END = re.compile(r"</w:p>")


def extract_docx(data: bytes) -> Extracted:
    errors: list[str] = []
    try:
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = normalize("\n".join(parts))
        return Extracted(text=text, fmt="docx")
    except Exception as exc:
        errors.append(f"python-docx: {type(exc).__name__}: {exc}")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
        return Extracted(text=_xml_to_text(xml), fmt="docx", notes=["raw XML fallback", *errors])
    except zipfile.BadZipFile as exc:
        errors.append(f"zip: BadZipFile: {exc}")
    except Exception as exc:
        errors.append(f"zip: {type(exc).__name__}: {exc}")
    recovered = _recover_truncated(data)
    if recovered:
        return Extracted(
            text=recovered, fmt="docx", truncated=True, notes=["recovered from truncated zip", *errors]
        )
    raise ExtractError("CORRUPT_DOCX", "; ".join(errors)[:300])


def _xml_to_text(xml: str) -> str:
    xml = _PARA_END.sub("\n", xml)
    return normalize(_TAG.sub("", xml))


def _recover_truncated(data: bytes) -> str:
    """A cut zip has no central directory. Try to inflate the document.xml local entry directly."""
    marker = b"word/document.xml"
    pos = data.find(marker)
    if pos < 0:
        return ""
    start = pos + len(marker)
    # local header extra field length sits just before the filename; be lenient and scan for deflate data
    for offset in range(start, min(start + 64, len(data))):
        try:
            d = zlib.decompressobj(-zlib.MAX_WBITS)
            raw = d.decompress(data[offset:])
            if raw.count(b"<w:") > 3:
                return _xml_to_text(raw.decode("utf-8", errors="replace"))
        except zlib.error:
            continue
    return ""
